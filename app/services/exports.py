"""Generate local DOCX and PDF meeting minutes."""
from io import BytesIO
from pathlib import Path


def to_docx(meeting, segments, actions):
    from docx import Document
    doc = Document()
    doc.add_heading(meeting["title"], 0)
    doc.add_paragraph(f"Дата: {meeting['meeting_date']}")
    doc.add_heading("Резюме", level=1)
    doc.add_paragraph(meeting["summary"])
    doc.add_heading("Транскрипт", level=1)
    for item in segments:
        doc.add_paragraph(f"[{int(item['start_seconds']//60):02d}:{int(item['start_seconds']%60):02d}] {item['display_name']}: {item['text']}")
    doc.add_heading("Поручения", level=1)
    for item in actions:
        doc.add_paragraph(f"{item['task']} — {item['responsible_person']}; срок: {item['deadline']}; статус: {item['status']}", style="List Bullet")
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def to_pdf(meeting, segments, actions):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from xml.sax.saxutils import escape
    # Use an installed Unicode font when available so Cyrillic/Kazakh text is preserved.
    font_paths = [Path("C:/Windows/Fonts/arial.ttf"), Path("C:/Windows/Fonts/ARIAL.TTF"),
                  Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
                  Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf")]
    font_path = next((path for path in font_paths if path.exists()), None)
    font_name = "Helvetica"
    if font_path:
        pdfmetrics.registerFont(TTFont("QazUnicode", str(font_path)))
        font_name = "QazUnicode"
    out = BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, title=meeting["title"])
    styles = getSampleStyleSheet()
    for style_name in ("Title", "Heading2", "Normal", "BodyText"):
        styles[style_name].fontName = font_name
    story = [Paragraph(escape(meeting["title"]), styles["Title"]),
             Paragraph("Дата: " + escape(meeting["meeting_date"]), styles["Normal"]), Spacer(1, 12),
             Paragraph("Резюме", styles["Heading2"]), Paragraph(escape(meeting["summary"]).replace("\n", "<br/>"), styles["BodyText"]),
             Spacer(1, 10), Paragraph("Транскрипт", styles["Heading2"])]
    for item in segments:
        stamp = f"{int(item['start_seconds']//60):02d}:{int(item['start_seconds']%60):02d}"
        story.append(Paragraph(f"[{stamp}] {escape(item['display_name'])}: {escape(item['text'])}", styles["BodyText"]))
    story.extend([Spacer(1, 10), Paragraph("Поручения", styles["Heading2"])])
    for item in actions:
        story.append(Paragraph(escape(f"{item['task']} — {item['responsible_person']}; срок: {item['deadline']}; статус: {item['status']}"), styles["BodyText"]))
    doc.build(story)
    return out.getvalue()
